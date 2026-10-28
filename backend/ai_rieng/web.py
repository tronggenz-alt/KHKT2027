"""API của AI Riêng.

- /admin/api/ai-rieng/* (CHỈ admin – dependency can_admin cho MỌI route, POST/DELETE thêm CSRF): tài liệu, kiến thức & quy tắc, bài mẫu,
  sửa sai + hàng chờ báo sai, Sân tập, Bộ đề kiểm tra, cài đặt, sao lưu.
- /api/bao-sai (người dùng đã đăng nhập): báo lời giải sai → hàng chờ; chỉ admin mới chuyển thành mục sửa sai.
Không route nào trả key AI; nội dung kho chỉ trả cho admin.
"""
from __future__ import annotations

import csv
import io
import json
import re
import time
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from .. import db, nhat_ky, pipeline
from ..auth import can_dang_nhap, gioi_han, kiem_csrf
from ..config import settings
from ..db import NguoiDung
from ..quan_tri_web import can_admin
from . import cham, kho, sao_luu, tach, xu_ly

TEP_TOI_DA = 30 * 1024 * 1024
SO_TEP_TOI_DA = 500
CSRF = [Depends(kiem_csrf)]

router = APIRouter(prefix="/admin/api/ai-rieng", dependencies=[Depends(can_admin), gioi_han("quan_tri")])
router_nguoi_dung = APIRouter()


def _loi(e: Exception, ma: int = 400) -> HTTPException:
    return HTTPException(ma, {"loi": str(e)})


def _ghi(request: Request, admin: NguoiDung, chi_tiet: str) -> None:
    nhat_ky.ghi("ai_rieng", request, admin, phuong_thuc="admin_web", chi_tiet=chi_tiet[:300])


# ================================================================ tổng quan + cài đặt
@router.get("/tong-quan")
def tong_quan():
    return {"thong_ke": kho.thong_ke(), "cai_dat": kho.cai_dat(), "nhung": xu_ly.trang_thai_nhung(), "chu_de": kho.CHU_DE,
            "embed_model": settings.embed_model, "co_openai": bool(settings.openai_api_key and settings.openai_model),
            "co_gemini": bool(settings.gemini_api_key), "lan_kiem_tra": kho.ds_lan_kiem_tra(10)}


@router.get("/cai-dat")
def xem_cai_dat():
    return {**kho.cai_dat(), "_mac_dinh_system": kho.SYSTEM_MAC_DINH}


@router.post("/cai-dat", dependencies=CSRF)
def luu_cai_dat(vao: dict, request: Request, admin: NguoiDung = Depends(can_admin)):
    try:
        cd = kho.dat_cai_dat(vao)
    except (ValueError, TypeError) as e:
        raise _loi(e) from e
    _ghi(request, admin, "Cài đặt: " + ", ".join(f"{k}={v}" for k, v in vao.items() if k != "system_prompt")
         + (" · sửa system prompt" if "system_prompt" in vao else ""))
    return cd


# ================================================================ tài liệu
def _luu_tep(tid: int, ten: str, du_lieu: bytes) -> str:
    duoi = re.sub(r"[^a-z0-9]", "", ten.lower().rsplit(".", 1)[-1])[:6] if "." in ten else "bin"
    p = xu_ly.thu_muc_tep() / f"{tid}.{duoi or 'bin'}"
    p.write_bytes(du_lieu)
    return str(p)


@router.get("/tai-lieu")
def ds_tai_lieu():
    return {"tai_lieu": kho.ds_tai_lieu(), "nhung": xu_ly.trang_thai_nhung()}


@router.post("/tai-lieu", dependencies=CSRF)
async def tai_len(request: Request, tep: list[UploadFile] = File(...), chu_de: str = Form(""), co_loi_giai: bool = Form(False),
                  duong_dan_tuong_doi: str = Form(""), admin: NguoiDung = Depends(can_admin)):
    if len(tep) > SO_TEP_TOI_DA:
        raise HTTPException(400, {"loi": f"Tối đa {SO_TEP_TOI_DA} file mỗi lần."})

    rel_paths: list[str] = []
    if duong_dan_tuong_doi:
        try:
            parsed = json.loads(duong_dan_tuong_doi)
            if isinstance(parsed, list):
                rel_paths = [str(x) for x in parsed]
        except Exception:
            rel_paths = []

    ket_qua, loi = [], []
    for i, f in enumerate(tep):
        du_lieu = await f.read(TEP_TOI_DA + 1)
        await f.close()
        ten_goc = (f.filename or "tai_lieu").split("/")[-1][:180]
        if len(du_lieu) > TEP_TOI_DA:
            loi.append(f"{ten_goc}: quá 30 MB")
            continue
        try:
            loai = tach.loai_tu_ten(ten_goc, du_lieu)
        except tach.LoiTaiLieu as e:
            loi.append(f"{ten_goc}: {e}")
            continue

        rel_str = rel_paths[i] if i < len(rel_paths) else ""
        chu_de_tep = chu_de.strip()
        ten_luu = ten_goc.rsplit(".", 1)[0] or ten_goc

        if rel_str:
            rel_p = Path(rel_str)
            if rel_p.parent and str(rel_p.parent) != ".":
                phan_con = rel_p.parent.as_posix()
                chu_de_tep = f"{chu_de_tep} / {phan_con}" if chu_de_tep else phan_con
                ten_luu = f"{phan_con}/{ten_luu}"[:180]

        tid = kho.tao_tai_lieu(ten_luu, "van_ban" if loai == "html" else loai, rel_str or ten_goc, "", chu_de_tep, co_loi_giai, len(du_lieu))
        kho.cap_nhat_tai_lieu(tid, tep=_luu_tep(tid, ten_goc, du_lieu))
        ket_qua.append(tid)

    if ket_qua:
        xu_ly.danh_thuc()
        _ghi(request, admin, f"Tải lên {len(ket_qua)} tài liệu" + (" (có lời giải)" if co_loi_giai else ""))
    return {"da_them": ket_qua, "loi": loi}


class XemTruocThuMucVao(BaseModel):
    duong_dan: str = Field(min_length=1, max_length=1000)
    de_quy: bool = True


@router.post("/tai-lieu/xem-truoc-thu-muc", dependencies=CSRF)
def xem_truoc_thu_muc(vao: XemTruocThuMucVao, admin: NguoiDung = Depends(can_admin)):
    thu_muc = Path(vao.duong_dan.strip()).expanduser().resolve()
    if not thu_muc.exists() or not thu_muc.is_dir():
        raise HTTPException(400, {"loi": f"Thư mục không tồn tại hoặc không phải thư mục: {thu_muc}"})

    hop_le, bo_qua = [], []
    mau = "**/*" if vao.de_quy else "*"
    for p in thu_muc.glob(mau):
        if not p.is_file():
            continue
        if p.name.startswith(".") or p.name.startswith("~$") or p.name in ("Thumbs.db", "desktop.ini"):
            continue
        try:
            sz = p.stat().st_size
            if sz > TEP_TOI_DA:
                bo_qua.append(f"{p.name} (quá 30 MB)")
                continue
            with open(p, "rb") as fp:
                dau = fp.read(1024)
            loai = tach.loai_tu_ten(p.name, dau)
            try:
                rel = str(p.relative_to(thu_muc))
            except ValueError:
                rel = p.name
            hop_le.append({"ten": p.name, "duong_dan_tuong_doi": rel, "loai": loai, "dung_luong": sz})
        except Exception:
            bo_qua.append(p.name)

    return {
        "ten_thu_muc": thu_muc.name,
        "duong_dan": str(thu_muc),
        "tong_hop_le": len(hop_le),
        "tong_bo_qua": len(bo_qua),
        "danh_sach_hop_le": hop_le[:100],
        "bo_qua": bo_qua[:20],
    }


class QuetThuMucVao(BaseModel):
    duong_dan: str = Field(min_length=1, max_length=1000)
    chu_de: str = Field("", max_length=200)
    co_loi_giai: bool = False
    de_quy: bool = True
    kem_thu_muc_con: bool = True


@router.post("/tai-lieu/quet-thu-muc", dependencies=CSRF)
def quet_thu_muc(vao: QuetThuMucVao, request: Request, admin: NguoiDung = Depends(can_admin)):
    thu_muc = Path(vao.duong_dan.strip()).expanduser().resolve()
    if not thu_muc.exists() or not thu_muc.is_dir():
        raise HTTPException(400, {"loi": f"Thư mục không tồn tại hoặc không phải thư mục: {thu_muc}"})

    ket_qua, loi, bo_qua = [], [], []
    mau = "**/*" if vao.de_quy else "*"
    danh_sach_tep = []

    for p in thu_muc.glob(mau):
        if not p.is_file():
            continue
        if p.name.startswith(".") or p.name.startswith("~$") or p.name in ("Thumbs.db", "desktop.ini"):
            continue
        danh_sach_tep.append(p)
        if len(danh_sach_tep) >= SO_TEP_TOI_DA:
            break

    for p in danh_sach_tep:
        try:
            sz = p.stat().st_size
        except OSError:
            bo_qua.append(f"{p.name}: không đọc được file")
            continue

        if sz > TEP_TOI_DA:
            loi.append(f"{p.name}: quá 30 MB")
            continue

        try:
            with open(p, "rb") as fp:
                dau = fp.read(2048)
            loai = tach.loai_tu_ten(p.name, dau)
        except tach.LoiTaiLieu:
            bo_qua.append(p.name)
            continue
        except Exception as e:
            loi.append(f"{p.name}: {e}")
            continue

        try:
            du_lieu = p.read_bytes()
        except Exception as e:
            loi.append(f"{p.name}: {e}")
            continue

        chu_de_tep = vao.chu_de.strip()
        rel_str = p.name
        try:
            rel = p.relative_to(thu_muc)
            rel_str = str(rel)
            if vao.kem_thu_muc_con and rel.parent and str(rel.parent) != ".":
                phan_con = rel.parent.as_posix()
                chu_de_tep = f"{chu_de_tep} / {phan_con}" if chu_de_tep else phan_con
        except ValueError:
            pass

        if not chu_de_tep:
            chu_de_tep = thu_muc.name

        ten_hien_thi = (rel_str.rsplit(".", 1)[0] if "." in rel_str else rel_str)[:180]
        tid = kho.tao_tai_lieu(
            ten_hien_thi,
            "van_ban" if loai == "html" else loai,
            str(p)[:500],
            "",
            chu_de_tep[:200],
            vao.co_loi_giai,
            len(du_lieu),
        )
        kho.cap_nhat_tai_lieu(tid, tep=_luu_tep(tid, p.name, du_lieu))
        ket_qua.append(tid)

    if ket_qua:
        xu_ly.danh_thuc()
        _ghi(request, admin, f"Quét thư mục «{thu_muc.name}»: nạp {len(ket_qua)} tài liệu" + (" (có lời giải)" if vao.co_loi_giai else ""))

    return {
        "da_them": ket_qua,
        "tong_hop_le": len(ket_qua),
        "bo_qua": bo_qua[:30],
        "so_bo_qua": len(bo_qua),
        "loi": loi[:30],
        "so_loi": len(loi),
    }


class LinkVao(BaseModel):
    url: str = Field(max_length=2000)
    ten: str = Field("", max_length=200)
    chu_de: str = Field("", max_length=200)
    co_loi_giai: bool = False


@router.post("/tai-lieu/link", dependencies=CSRF)
def them_link(vao: LinkVao, request: Request, admin: NguoiDung = Depends(can_admin)):
    url = vao.url.strip()
    if not re.match(r"https?://[^\s/]+", url):
        raise HTTPException(400, {"loi": "Link phải bắt đầu bằng http:// hoặc https://"})
    tid = kho.tao_tai_lieu(vao.ten.strip() or url[:120], "link", url, "", vao.chu_de, vao.co_loi_giai)
    xu_ly.danh_thuc()
    _ghi(request, admin, f"Thêm link tài liệu {url[:120]}")
    return {"id": tid}


class VanBanVao(BaseModel):
    ten: str = Field(max_length=200)
    noi_dung: str = Field(max_length=2_000_000)
    chu_de: str = Field("", max_length=200)
    co_loi_giai: bool = False


@router.post("/tai-lieu/van-ban", dependencies=CSRF)
def them_van_ban(vao: VanBanVao, request: Request, admin: NguoiDung = Depends(can_admin)):
    if not vao.ten.strip() or len(vao.noi_dung.strip()) < 20:
        raise HTTPException(400, {"loi": "Cần tên và nội dung (ít nhất 20 ký tự)."})
    du_lieu = vao.noi_dung.encode("utf-8")
    tid = kho.tao_tai_lieu(vao.ten.strip(), "van_ban", "dán vào", "", vao.chu_de, vao.co_loi_giai, len(du_lieu))
    kho.cap_nhat_tai_lieu(tid, tep=_luu_tep(tid, "van_ban.txt", du_lieu))
    xu_ly.danh_thuc()
    _ghi(request, admin, f"Dán văn bản «{vao.ten.strip()[:80]}»")
    return {"id": tid}


@router.get("/tai-lieu/{tid}/doan")
def xem_doan(tid: int):
    tl = kho.lay_tai_lieu(tid)
    if tl is None:
        raise HTTPException(404, {"loi": "Không có tài liệu này."})
    return {"tai_lieu": tl, "doan": kho.ds_doan(tid)}


class SuaTaiLieuVao(BaseModel):
    ten: str | None = Field(None, max_length=200)
    chu_de: str | None = Field(None, max_length=200)
    bat: bool | None = None


@router.post("/tai-lieu/{tid}", dependencies=CSRF)
def sua_tai_lieu(tid: int, vao: SuaTaiLieuVao):
    if kho.lay_tai_lieu(tid) is None:
        raise HTTPException(404, {"loi": "Không có tài liệu này."})
    kho.cap_nhat_tai_lieu(tid, **{k: (int(v) if isinstance(v, bool) else v) for k, v in vao.model_dump().items() if v is not None})
    return kho.lay_tai_lieu(tid)


@router.post("/tai-lieu/{tid}/xu-ly-lai", dependencies=CSRF)
def xu_ly_lai(tid: int):
    tl = kho.lay_tai_lieu(tid)
    if tl is None:
        raise HTTPException(404, {"loi": "Không có tài liệu này."})
    if tl["loai"] != "link" and not (tl["tep"] and Path(tl["tep"]).is_file()):
        raise HTTPException(400, {"loi": "Không còn file gốc (tài liệu nhập từ bản sao lưu) – chỉ sửa được từng đoạn."})
    kho.cap_nhat_tai_lieu(tid, trang_thai="cho", tien_do="Xếp hàng xử lý lại", loi="")
    xu_ly.danh_thuc()
    return {"ok": True}


@router.delete("/tai-lieu/{tid}", dependencies=CSRF)
def xoa_tai_lieu(tid: int, request: Request, admin: NguoiDung = Depends(can_admin)):
    tl = kho.xoa_tai_lieu(tid)
    if tl is None:
        raise HTTPException(404, {"loi": "Không có tài liệu này."})
    if tl["tep"]:
        Path(tl["tep"]).unlink(missing_ok=True)
    _ghi(request, admin, f"Xoá tài liệu «{tl['ten'][:80]}»")
    return {"ok": True}


class DoanVao(BaseModel):
    noi_dung: str = Field(max_length=20000)


@router.post("/doan/{did}", dependencies=CSRF)
def sua_doan(did: int, vao: DoanVao):
    if not kho.sua_doan(did, vao.noi_dung):
        raise HTTPException(400, {"loi": "Không sửa được đoạn này (không tồn tại hoặc nội dung trống)."})
    return {"ok": True}


@router.delete("/doan/{did}", dependencies=CSRF)
def xoa_doan(did: int):
    if not kho.xoa_doan(did):
        raise HTTPException(404, {"loi": "Không có đoạn này."})
    return {"ok": True}


# ================================================================ kiến thức & quy tắc
class KienThucVao(BaseModel):
    tieu_de: str = Field(max_length=300)
    noi_dung: str = Field(max_length=8000)
    chu_de: str = Field("", max_length=200)
    uu_tien: int = Field(3, ge=1, le=5)
    bat: bool = True


@router.get("/kien-thuc")
def ds_kien_thuc(chu_de: str = "", q: str = ""):
    return {"kien_thuc": kho.ds_kien_thuc(chu_de, q[:200])}


@router.post("/kien-thuc", dependencies=CSRF)
def them_kien_thuc(vao: KienThucVao):
    try:
        return {"id": kho.tao_kien_thuc(**vao.model_dump())}
    except ValueError as e:
        raise _loi(e) from e


@router.post("/kien-thuc/nhap", dependencies=CSRF)
async def nhap_kien_thuc(request: Request, tep: UploadFile = File(...), admin: NguoiDung = Depends(can_admin)):
    du_lieu = await tep.read(5 * 1024 * 1024 + 1)
    if len(du_lieu) > 5 * 1024 * 1024:
        raise HTTPException(400, {"loi": "File tối đa 5 MB."})
    try:
        hang = doc_bang(tep.filename or "", du_lieu)
    except ValueError as e:
        raise _loi(e) from e
    them, loi = 0, []
    for i, h in enumerate(hang, 2):
        try:
            kho.tao_kien_thuc(h.get("tieu_de", ""), h.get("noi_dung", ""), h.get("chu_de", ""), int(h.get("uu_tien") or 3))
            them += 1
        except (ValueError, TypeError) as e:
            loi.append(f"Dòng {i}: {e}")
    _ghi(request, admin, f"Nhập {them} quy tắc từ {tep.filename}")
    return {"da_nhap": them, "loi": loi[:50]}


@router.post("/kien-thuc/{kid}", dependencies=CSRF)
def sua_kien_thuc(kid: int, vao: KienThucVao):
    try:
        if not kho.sua_kien_thuc(kid, **vao.model_dump()):
            raise HTTPException(404, {"loi": "Không có mục này."})
    except ValueError as e:
        raise _loi(e) from e
    return {"ok": True}


@router.delete("/kien-thuc/{kid}", dependencies=CSRF)
def xoa_kien_thuc(kid: int):
    if not kho.xoa_kien_thuc(kid):
        raise HTTPException(404, {"loi": "Không có mục này."})
    return {"ok": True}


_COT = {"tieu_de": ("tieu_de", "tiêu đề", "tieu de", "title", "ten", "tên"), "noi_dung": ("noi_dung", "nội dung", "noi dung", "content"),
        "chu_de": ("chu_de", "chủ đề", "chu de", "topic"), "uu_tien": ("uu_tien", "độ ưu tiên", "ưu tiên", "do uu tien", "priority"),
        "de": ("de", "đề", "đề bài", "de bai", "chuoi", "chuỗi"), "dap_an": ("dap_an", "đáp án", "dap an", "answer")}


def doc_bang(ten: str, du_lieu: bytes) -> list[dict]:
    """CSV (UTF-8, dấu , hoặc ;) hoặc Excel .xlsx → [{cột chuẩn: giá trị}] theo dòng tiêu đề."""
    if ten.lower().endswith(".xlsx") or du_lieu[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            ws = load_workbook(io.BytesIO(du_lieu), read_only=True, data_only=True).active
            dong = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]
        except Exception as e:  # noqa: BLE001
            raise ValueError("Không đọc được file Excel (.xlsx).") from e
    else:
        text = du_lieu.decode("utf-8-sig", "ignore")
        try:
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        dong = list(csv.reader(io.StringIO(text), dialect))
    dong = [r for r in dong if any(str(x).strip() for x in r)]
    if len(dong) < 2:
        raise ValueError("File cần dòng tiêu đề + ít nhất một dòng dữ liệu.")
    tieu_de = [str(x).strip().lower() for x in dong[0]]
    chi_so = {chuan: next((i for i, t in enumerate(tieu_de) if t in ten_cot), None) for chuan, ten_cot in _COT.items()}
    return [{k: str(r[i]).strip() for k, i in chi_so.items() if i is not None and i < len(r)} for r in dong[1:]]


# ================================================================ bài mẫu (thư viện lời giải)
class BaiMauVao(BaseModel):
    de: str = Field(max_length=4000)
    loi_giai: dict
    chu_de: str = Field("", max_length=200)
    trang_thai: str = Field("duyet", pattern="^(duyet|cho_duyet|sai)$")
    bat: bool = True


def _kiem_rdkit(de: str, lg: dict) -> list[str]:
    from .. import giai_nhanh
    from ..schemas import LoiGiaiGon

    try:
        gon = LoiGiaiGon.model_validate(lg)
    except ValueError as e:
        raise HTTPException(400, {"loi": f"Lời giải không đúng cấu trúc: {e}"[:300]}) from e
    return giai_nhanh.kiem_tra(gon, de)[0]


@router.get("/bai-mau")
def ds_bai_mau(trang_thai: str = "", q: str = "", chu_de: str = ""):
    return {"bai_mau": kho.ds_bai_mau(trang_thai, q[:200], chu_de)}


@router.post("/bai-mau/kiem-tra", dependencies=CSRF)
def kiem_tra_bai_mau(vao: BaiMauVao):
    loi = _kiem_rdkit(vao.de, vao.loi_giai)
    return {"rdkit_ok": not loi, "loi": loi}


@router.post("/bai-mau", dependencies=CSRF)
def them_bai_mau(vao: BaiMauVao):
    loi = _kiem_rdkit(vao.de, vao.loi_giai)
    try:
        return {"id": kho.tao_bai_mau(vao.de, vao.loi_giai, vao.chu_de, vao.trang_thai, "tay", rdkit_ok=not loi, bat=vao.bat), "loi_rdkit": loi}
    except ValueError as e:
        raise _loi(e) from e


class DuyetVao(BaseModel):
    ids: list[int] = Field(max_length=1000)
    trang_thai: str = Field("duyet", pattern="^(duyet|cho_duyet|sai)$")


@router.post("/bai-mau/duyet", dependencies=CSRF)
def duyet(vao: DuyetVao):
    return {"so": kho.duyet_bai_mau(vao.ids, vao.trang_thai)}


@router.post("/bai-mau/{bid}", dependencies=CSRF)
def sua_bai_mau(bid: int, vao: BaiMauVao):
    loi = _kiem_rdkit(vao.de, vao.loi_giai)
    try:
        if not kho.sua_bai_mau(bid, vao.de, vao.loi_giai, vao.chu_de, vao.trang_thai, not loi, vao.bat):
            raise HTTPException(404, {"loi": "Không có bài mẫu này."})
    except ValueError as e:
        raise _loi(e) from e
    return {"ok": True, "loi_rdkit": loi}


@router.delete("/bai-mau/{bid}", dependencies=CSRF)
def xoa_bai_mau(bid: int):
    if not kho.xoa_bai_mau(bid):
        raise HTTPException(404, {"loi": "Không có bài mẫu này."})
    return {"ok": True}


# ================================================================ sửa sai + hàng chờ báo sai
class SuaSaiVao(BaseModel):
    de: str = Field(max_length=4000)
    dung: str = Field(max_length=6000)
    ly_do: str = Field("", max_length=3000)
    sai: str = Field("", max_length=6000)
    chu_de: str = Field("", max_length=200)
    bat: bool = True


@router.get("/sua-sai")
def ds_sua_sai(q: str = ""):
    return {"sua_sai": kho.ds_sua_sai(q[:200]), "bao_sai_cho": kho.so_bao_sai_cho()}


@router.post("/sua-sai", dependencies=CSRF)
def them_sua_sai(vao: SuaSaiVao):
    try:
        return {"id": kho.tao_sua_sai(vao.de, vao.dung, vao.ly_do, vao.sai, vao.chu_de)}
    except ValueError as e:
        raise _loi(e) from e


@router.post("/sua-sai/{sid}", dependencies=CSRF)
def sua_sua_sai(sid: int, vao: SuaSaiVao):
    try:
        if not kho.sua_sua_sai(sid, **vao.model_dump()):
            raise HTTPException(404, {"loi": "Không có mục này."})
    except ValueError as e:
        raise _loi(e) from e
    return {"ok": True}


@router.delete("/sua-sai/{sid}", dependencies=CSRF)
def xoa_sua_sai(sid: int):
    if not kho.xoa_sua_sai(sid):
        raise HTTPException(404, {"loi": "Không có mục này."})
    return {"ok": True}


@router.get("/bao-sai")
def ds_bao_sai(trang_thai: str = "cho"):
    return {"bao_sai": kho.ds_bao_sai(trang_thai if trang_thai in ("cho", "da_sua", "bo_qua") else "")}


class ChuyenBaoSaiVao(BaseModel):
    dung: str = Field(max_length=6000)
    ly_do: str = Field("", max_length=3000)
    chu_de: str = Field("", max_length=200)


@router.post("/bao-sai/{bid}/chuyen", dependencies=CSRF)
def chuyen_bao_sai(bid: int, vao: ChuyenBaoSaiVao, request: Request, admin: NguoiDung = Depends(can_admin)):
    b = kho.lay_bao_sai(bid)
    if b is None:
        raise HTTPException(404, {"loi": "Không có báo sai này."})
    try:
        sid = kho.tao_sua_sai(b["de"], vao.dung, vao.ly_do, tom_tat_sai(b["loi_giai"], b["mo_hinh"], b["ghi_chu"]), vao.chu_de, bao_sai_id=bid)
    except ValueError as e:
        raise _loi(e) from e
    kho.xu_ly_bao_sai(bid, "da_sua", sid)
    _ghi(request, admin, f"Duyệt báo sai #{bid} → sửa sai #{sid}")
    return {"sua_sai_id": sid}


@router.post("/bao-sai/{bid}/bo-qua", dependencies=CSRF)
def bo_qua_bao_sai(bid: int):
    if not kho.xu_ly_bao_sai(bid, "bo_qua"):
        raise HTTPException(404, {"loi": "Không có báo sai này."})
    return {"ok": True}


def tom_tat_sai(gon: dict, mo_hinh: str = "", ghi_chu: str = "") -> str:
    """Lời giải bị sai (JSON gọn) → vài dòng ghi vào mục sửa sai (để AI biết mình từng sai thế nào)."""
    chu = kho.tom_tat_loi_giai(gon or {})
    return ((f"[{mo_hinh}] " if mo_hinh else "") + chu + (f"\nNgười dùng ghi: {ghi_chu}" if ghi_chu else ""))[:6000]


@router.get("/muc/{loai}/{mid}")
def xem_muc(loai: str, mid: int):
    d = kho.chi_tiet_muc(loai, mid)
    if d is None:
        raise HTTPException(404, {"loi": "Không còn mục này (có thể đã bị xoá)."})
    return {"loai": loai, "muc": d}


# ================================================================ Sân tập
class SanTapVao(BaseModel):
    de: str = Field(min_length=3, max_length=4000)
    so_sanh: bool = False
    bo_qua_thu_vien: bool = True
    dung_claude: bool = False


@router.post("/san-tap", dependencies=CSRF)
def san_tap(vao: SanTapVao, admin: NguoiDung = Depends(can_admin)):
    from .. import main

    try:
        out = {"ai_rieng": main.bat_dau_giai_san_tap(admin, vao.de, "ai_rieng", {"bo_qua_thu_vien": vao.bo_qua_thu_vien, "san_tap": True},
                                                     vao.dung_claude)}
        if vao.so_sanh:
            out["ai_thuong"] = main.bat_dau_giai_san_tap(admin, vao.de, "", {}, vao.dung_claude)
    except Exception as e:  # noqa: BLE001 – vd hết key AI
        raise HTTPException(400, {"loi": str(e)[:300]}) from e
    return out


def _ket_qua_job(job_id: str, u: NguoiDung) -> dict:
    job = pipeline.JOBS.get(job_id)
    if job is not None and job.user_id == u.id and job.ket_qua:
        return job.ket_qua
    da_luu = db.ket_qua_da_luu(job_id, u.id) if re.fullmatch(r"\d{8}-\d{6}-[0-9a-f]{6,32}", job_id) else None
    if not da_luu or not da_luu.get("ket_qua"):
        raise HTTPException(404, {"loi": "Không tìm thấy kết quả bài giải này."})
    return da_luu["ket_qua"]


class DungVao(BaseModel):
    job_id: str = Field(max_length=64)
    chu_de: str = Field("", max_length=200)


@router.post("/san-tap/dung", dependencies=CSRF)
def san_tap_dung(vao: DungVao, admin: NguoiDung = Depends(can_admin)):
    """"Đúng" → lưu lời giải thành bài mẫu ĐÃ DUYỆT (thay bản chờ duyệt cùng đề nếu có)."""
    kq = _ket_qua_job(vao.job_id, admin)
    de = (kq.get("ai_rieng") or {}).get("de") or kq.get("de_bai", "")
    gon = kq.get("gon")
    if not gon or not de:
        raise HTTPException(400, {"loi": "Bài này không có lời giải dạng chuỗi để lưu."})
    cung_de = [b for b in kho.ds_bai_mau(q="") if b["de_khoa"] == kho.khoa_de(de) and b["trang_thai"] != "sai"]
    loi = _kiem_rdkit(de, gon)
    if cung_de:
        b = cung_de[0]
        kho.sua_bai_mau(b["id"], de, gon, vao.chu_de or b["chu_de"], "duyet", not loi)
        return {"id": b["id"], "cap_nhat": True}
    return {"id": kho.tao_bai_mau(de, gon, vao.chu_de, "duyet", "san_tap", rdkit_ok=not loi, mo_hinh=kq.get("mo_hinh", "")), "cap_nhat": False}


class SaiVao(BaseModel):
    job_id: str = Field(max_length=64)
    dung: str = Field(max_length=6000)
    ly_do: str = Field("", max_length=3000)
    chu_de: str = Field("", max_length=200)


@router.post("/san-tap/sai", dependencies=CSRF)
def san_tap_sai(vao: SaiVao, admin: NguoiDung = Depends(can_admin)):
    """"Sai – sửa" → mục sửa sai (ưu tiên cao nhất) + lời giải cùng đề trong thư viện không được dùng lại."""
    kq = _ket_qua_job(vao.job_id, admin)
    de = (kq.get("ai_rieng") or {}).get("de") or kq.get("de_bai", "")
    try:
        sid = kho.tao_sua_sai(de, vao.dung, vao.ly_do, tom_tat_sai(kq.get("gon") or {}, kq.get("mo_hinh", "")), vao.chu_de)
    except ValueError as e:
        raise _loi(e) from e
    return {"id": sid}


# ================================================================ bộ đề kiểm tra
class CauHoiVao(BaseModel):
    de: str = Field(max_length=4000)
    dap_an: list[dict] = Field(max_length=40)
    chu_de: str = Field("", max_length=200)
    bat: bool = True


@router.get("/bo-de")
def ds_bo_de():
    return {"bo_de": kho.ds_cau_hoi(), "lan": kho.ds_lan_kiem_tra(50)}


@router.post("/bo-de", dependencies=CSRF)
def them_cau_hoi(vao: CauHoiVao):
    try:
        return {"id": kho.tao_cau_hoi(vao.de, vao.dap_an, vao.chu_de, vao.bat)}
    except ValueError as e:
        raise _loi(e) from e


def doc_dap_an(text: str) -> list[dict]:
    """"A: CH3CHO; B: CH3COOH" (hoặc A = …) → [{ki_hieu, chat}]."""
    out = []
    for phan in re.split(r"[;\n]", text or ""):
        m = re.match(r"\s*([^:=]{1,20}?)\s*[:=]\s*(.+?)\s*$", phan)
        if m:
            out.append({"ki_hieu": m.group(1).strip(), "chat": m.group(2).strip()})
    return out


@router.post("/bo-de/nhap", dependencies=CSRF)
async def nhap_bo_de(tep: UploadFile = File(...)):
    du_lieu = await tep.read(5 * 1024 * 1024 + 1)
    if len(du_lieu) > 5 * 1024 * 1024:
        raise HTTPException(400, {"loi": "File tối đa 5 MB."})
    try:
        hang = doc_bang(tep.filename or "", du_lieu)
    except ValueError as e:
        raise _loi(e) from e
    them, loi = 0, []
    for i, h in enumerate(hang, 2):
        try:
            kho.tao_cau_hoi(h.get("de", ""), doc_dap_an(h.get("dap_an", "")), h.get("chu_de", ""))
            them += 1
        except ValueError as e:
            loi.append(f"Dòng {i}: {e}")
    return {"da_nhap": them, "loi": loi[:50]}


@router.post("/bo-de/{cid}", dependencies=CSRF)
def sua_cau_hoi(cid: int, vao: CauHoiVao):
    try:
        if not kho.sua_cau_hoi(cid, vao.de, vao.dap_an, vao.chu_de, vao.bat):
            raise HTTPException(404, {"loi": "Không có câu này."})
    except ValueError as e:
        raise _loi(e) from e
    return {"ok": True}


@router.delete("/bo-de/{cid}", dependencies=CSRF)
def xoa_cau_hoi(cid: int):
    if not kho.xoa_cau_hoi(cid):
        raise HTTPException(404, {"loi": "Không có câu này."})
    return {"ok": True}


@router.post("/kiem-tra/chay", dependencies=CSRF)
def chay_kiem_tra(request: Request, admin: NguoiDung = Depends(can_admin)):
    if not (settings.gemini_api_key or settings.api_key or settings.openai_api_key):
        raise HTTPException(400, {"loi": "Máy chủ chưa có key AI nào để chạy kiểm tra."})
    try:
        lid = cham.bat_dau()
    except ValueError as e:
        raise _loi(e) from e
    _ghi(request, admin, f"Chạy bộ đề kiểm tra (lần #{lid})")
    return {"id": lid}


@router.get("/kiem-tra/{lid}")
def xem_kiem_tra(lid: int):
    d = kho.lay_lan_kiem_tra(lid)
    if d is None:
        raise HTTPException(404, {"loi": "Không có lần kiểm tra này."})
    return {**d, "theo_chu_de": cham.theo_chu_de(d["chi_tiet"])}


# ================================================================ sao lưu
@router.get("/sao-luu")
def tai_sao_luu(request: Request, admin: NguoiDung = Depends(can_admin)):
    du_lieu = sao_luu.xuat()
    _ghi(request, admin, "Tải bản sao lưu kho AI Riêng")
    ten = f"ai-rieng-{time.strftime('%Y%m%d-%H%M')}.json"
    return Response(json.dumps(du_lieu, ensure_ascii=False, indent=1), media_type="application/json",
                    headers={"Content-Disposition": f'attachment; filename="{ten}"', "Cache-Control": "no-store"})


@router.post("/phuc-hoi", dependencies=CSRF)
async def phuc_hoi(request: Request, tep: UploadFile = File(...), thay_the: bool = Form(False), nhap_cai_dat: bool = Form(False),
                   admin: NguoiDung = Depends(can_admin)):
    du_lieu = await tep.read(200 * 1024 * 1024 + 1)
    if len(du_lieu) > 200 * 1024 * 1024:
        raise HTTPException(400, {"loi": "File sao lưu tối đa 200 MB."})
    try:
        dem = sao_luu.nhap(json.loads(du_lieu.decode("utf-8-sig")), thay_the, nhap_cai_dat)
    except (ValueError, KeyError, TypeError) as e:
        raise _loi(e if isinstance(e, ValueError) else ValueError(f"File sao lưu thiếu trường {e}.")) from e
    xu_ly.danh_thuc()
    _ghi(request, admin, ("Thay toàn bộ" if thay_the else "Gộp") + f" kho từ bản sao lưu: {dem}")
    return dem


# ================================================================ người dùng: báo sai
class BaoSaiVao(BaseModel):
    job_id: str = Field(max_length=64)
    ghi_chu: str = Field("", max_length=2000)


@router_nguoi_dung.post("/api/bao-sai", dependencies=[Depends(kiem_csrf), gioi_han("bao_sai")])
def bao_sai(vao: BaoSaiVao, request: Request, u: NguoiDung = Depends(can_dang_nhap)):
    """Người dùng báo lời giải sai → hàng chờ admin (người dùng KHÔNG dạy trực tiếp được)."""
    if not re.fullmatch(r"\d{8}-\d{6}-[0-9a-f]{6,32}", vao.job_id):
        raise HTTPException(404, {"loi": "Không tìm thấy bài giải này."})
    job = pipeline.JOBS.get(vao.job_id)
    chu = job.user_id if job is not None else db.chu_bai_giai(vao.job_id)
    if chu != u.id:
        raise HTTPException(404, {"loi": "Không tìm thấy bài giải này."})
    kq = job.ket_qua if job is not None and job.ket_qua else (db.ket_qua_da_luu(vao.job_id, u.id) or {}).get("ket_qua")
    if not kq or kq.get("demo"):
        raise HTTPException(400, {"loi": "Chỉ báo sai được lời giải AI đã giải xong."})
    bid = kho.tao_bao_sai(u.id, vao.job_id, (kq.get("ai_rieng") or {}).get("de") or kq.get("de_bai", ""), kq.get("mo_hinh", ""),
                          kq.get("gon"), vao.ghi_chu.strip())
    if bid is None:
        raise HTTPException(409, {"loi": "Bạn đã báo sai bài này rồi – admin sẽ xem."})
    nhat_ky.ghi("bao_sai_ai", request, u, phuong_thuc="web", chi_tiet=f"Bài {vao.job_id}: {vao.ghi_chu.strip()[:200]}")
    return {"ok": True, "id": bid}
